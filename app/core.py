"""Dominio de la demo: boletas sintéticas, políticas, estado y auditoría."""
from __future__ import annotations

import hashlib
import os
import re
import sqlite3
import subprocess
import tempfile
import uuid
from datetime import date, datetime, timezone
from pathlib import Path

from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DATA = Path(os.environ.get("DEMO_DATA", ROOT / "data")).resolve()
DB = DATA / "rendiciones.sqlite3"
BOOK = DATA / "rendiciones.xlsx"
HEADERS = ["ID", "Empleado", "Proveedor", "Folio", "Fecha", "Categoría", "Monto CLP", "Aprobación", "Boleta", "Correo"]
MAX_AMOUNT = int(os.environ.get("AUTO_APPROVE_MAX_CLP", "50000"))


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect() -> sqlite3.Connection:
    DATA.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB, timeout=20)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA busy_timeout=20000")
    return con


def init() -> None:
    (DATA / "uploads").mkdir(parents=True, exist_ok=True)
    (DATA / "runs").mkdir(parents=True, exist_ok=True)
    with connect() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS claims (
            id TEXT PRIMARY KEY, employee TEXT NOT NULL, category TEXT NOT NULL,
            merchant TEXT NOT NULL, receipt_number TEXT NOT NULL, receipt_date TEXT NOT NULL,
            amount INTEGER NOT NULL, filename TEXT NOT NULL, original_name TEXT NOT NULL,
            sha256 TEXT NOT NULL, ocr_text TEXT NOT NULL, state TEXT NOT NULL,
            reason TEXT NOT NULL, email_to TEXT NOT NULL, error TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
            UNIQUE(merchant, receipt_number)
        );
        CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY AUTOINCREMENT, claim_id TEXT NOT NULL,
            at TEXT NOT NULL, step TEXT NOT NULL, detail TEXT NOT NULL,
            screenshot TEXT NOT NULL DEFAULT ''
        );
        """)
    if not BOOK.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Rendiciones"
        ws.append(HEADERS)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:J1"
        for i, width in enumerate((39, 24, 31, 19, 17, 18, 18, 23, 29, 24), 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        for cell in ws[1]:
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 27
        wb.save(BOOK)


def get_claim(claim_id: str) -> dict | None:
    with connect() as con:
        row = con.execute("SELECT * FROM claims WHERE id=?", (claim_id,)).fetchone()
    return dict(row) if row else None


def list_claims() -> list[dict]:
    with connect() as con:
        rows = con.execute("SELECT * FROM claims ORDER BY created_at DESC LIMIT 100").fetchall()
    return [dict(r) for r in rows]


def list_events(claim_id: str) -> list[dict]:
    with connect() as con:
        rows = con.execute("SELECT at,step,detail,screenshot FROM events WHERE claim_id=? ORDER BY id", (claim_id,)).fetchall()
    return [dict(r) for r in rows]


def event(claim_id: str, step: str, detail: str, screenshot: str = "") -> None:
    with connect() as con:
        con.execute("INSERT INTO events (claim_id,at,step,detail,screenshot) VALUES (?,?,?,?,?)",
                    (claim_id, now(), step, detail, screenshot))


def transition(claim_id: str, from_states: tuple[str, ...], to_state: str, **fields) -> bool:
    if not from_states:
        raise ValueError("Se requiere al menos un estado de origen")
    keys = ["state", "updated_at", *fields]
    values = [to_state, now(), *fields.values()]
    marks = ",".join(f"{k}=?" for k in keys)
    where = ",".join("?" for _ in from_states)
    with connect() as con:
        result = con.execute(f"UPDATE claims SET {marks} WHERE id=? AND state IN ({where})",
                             [*values, claim_id, *from_states])
    return result.rowcount == 1


def extract(path: Path) -> tuple[dict, str]:
    """OCR local. Nunca inventa campos ausentes; obliga a revisión humana."""
    source = path
    with tempfile.TemporaryDirectory() as tmp:
        if path.suffix.lower() == ".pdf":
            source = Path(tmp) / "pagina.png"
            subprocess.run(["pdftoppm", "-f", "1", "-l", "1", "-singlefile", "-r", "220", "-png", str(path), str(source.with_suffix(""))],
                           check=True, timeout=25, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        else:
            with Image.open(path) as image:
                image.verify()
        try:
            cmd = ["tesseract", str(source), "stdout", "-l", "spa+eng", "--psm", "6"]
            text = subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=25).stdout
        except (subprocess.CalledProcessError, FileNotFoundError):
            text = subprocess.run(["tesseract", str(source), "stdout", "--psm", "6"],
                                  check=True, capture_output=True, text=True, timeout=25).stdout
    text = text.strip()

    def field(patterns: list[str]) -> str:
        for pattern in patterns:
            found = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
            if found:
                return found.group(1).strip()
        return ""

    merchant = field([r"^(?:EMISOR|COMERCIO|PROVEEDOR)\s*[:\-]?\s*(.+)$"])
    number = field([r"^(?:FOLIO|N[º°o.]|NUMERO)\s*[:\-#]?\s*([A-Z0-9-]{3,25})\s*$"])
    issued = field([r"^FECHA\s*[:\-]?\s*(\d{4}[-/]\d{2}[-/]\d{2}|\d{2}[-/]\d{2}[-/]\d{4})"])
    amount = field([r"^TOTAL\s*[:\-]?\s*\$?\s*([\d.,]+)"])
    amount_int = int(re.sub(r"\D", "", amount)) if amount else 0
    if issued:
        parts = issued.replace("/", "-").split("-")
        issued = "-".join(parts if len(parts[0]) == 4 else reversed(parts))
    return {"merchant": merchant[:100], "receipt_number": number[:25],
            "receipt_date": issued, "amount": amount_int}, text


def validate(data: dict, text: str) -> list[str]:
    reasons = []
    for key, label in (("merchant", "proveedor"), ("receipt_number", "folio"), ("receipt_date", "fecha"), ("amount", "monto")):
        if not data[key]:
            reasons.append(f"OCR: falta {label}")
    try:
        issued = date.fromisoformat(data["receipt_date"])
        if issued > date.today() or (date.today() - issued).days > 365:
            reasons.append("Fecha fuera de política (futura o mayor a 365 días)")
    except ValueError:
        reasons.append("Fecha ilegible o inválida")
    if data["amount"] > MAX_AMOUNT:
        reasons.append(f"Monto mayor que el máximo de autoaprobación ({MAX_AMOUNT:,} CLP)")
    if data["amount"] <= 0:
        reasons.append("Monto no válido")
    if len(text) < 25:
        reasons.append("Documento ilegible: revisión obligatoria")
    return list(dict.fromkeys(reasons))


def create_claim(file, employee: str, category: str) -> dict:
    init()
    extension = Path(file.filename or "").suffix.lower()
    if extension not in (".png", ".jpg", ".jpeg", ".pdf"):
        raise ValueError("Formato permitido: PNG, JPG o PDF")
    employee = employee.strip()[:80]
    category = category.strip()[:40]
    if not employee or not category:
        raise ValueError("Empleado y categoría son obligatorios")
    content = file.read()
    if not content or len(content) > 5 * 1024 * 1024:
        raise ValueError("Documento vacío o mayor de 5 MB")
    claim_id = str(uuid.uuid4())
    filename = claim_id + extension
    path = DATA / "uploads" / filename
    path.write_bytes(content)
    try:
        data, ocr_text = extract(path)
        reasons = validate(data, ocr_text)
        state = "pendiente_aprobacion" if reasons else "listo"
        with connect() as con:
            con.execute("""INSERT INTO claims
                (id,employee,category,merchant,receipt_number,receipt_date,amount,filename,original_name,sha256,
                 ocr_text,state,reason,email_to,created_at,updated_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (claim_id, employee, category, data["merchant"], data["receipt_number"], data["receipt_date"],
                 data["amount"], filename, Path(file.filename or "boleta").name[:120], hashlib.sha256(content).hexdigest(),
                 ocr_text, state, "; ".join(reasons), "demo@ejemplo.local", now(), now()))
    except sqlite3.IntegrityError as exc:
        path.unlink(missing_ok=True)
        raise ValueError("Boleta duplicada: proveedor y folio ya registrados") from exc
    except Exception:
        path.unlink(missing_ok=True)
        raise
    event(claim_id, "OCR", "Campos extraídos con Tesseract local; los datos requieren revisión si hay excepciones")
    event(claim_id, "Política", "; ".join(reasons) if reasons else f"Autoaprobación: monto ≤ {MAX_AMOUNT:,} CLP")
    return get_claim(claim_id)


def book_row(claim_id: str) -> int | None:
    if not BOOK.exists():
        return None
    wb = load_workbook(BOOK, read_only=True)
    try:
        for i, cells in enumerate(wb.active.iter_rows(min_row=2, max_col=1, values_only=True), start=2):
            if str(cells[0]) == claim_id:
                return i
    finally:
        wb.close()
    return None
